"""Create the explanatory guide and English submission with numbered IEEE references."""
from pathlib import Path
from datetime import datetime
import json, re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'/'entrega'; OUT.mkdir(parents=True,exist_ok=True)
AUTHORS=['José Luis Paladines Sánchez','John Maximiliano Carreño Robalino','Julinho Jesús Plua Mendosa']
REFS=[
 '[1] IEEE-HKN Nu Zeta and Nu Eta, “Reglas del Hackathon Hack2Impact 2026,” competition brief supplied by the team, 2026.',
 '[2] Instituto Nacional de Estadística y Censos, “Marco Geoestadístico 2022,” FeatureServer layer 7. Accessed: Oct. 25, 2026. [Online]. Available: https://idgn.ecuadorencifras.gob.ec/server/rest/services/Hosted/Marco_Geoestadistico_2022/FeatureServer/7',
 '[3] Municipalidad de Guayaquil, “Sectores de nomenclatura,” Sistema de Información Local, updated Nov. 28, 2025. Accessed: Oct. 25, 2026. [Online]. Available: https://sil.guayaquil.gob.ec/territorial/sectores-de-nomenclatura/',
 '[4] Open-Meteo, “Single Runs API,” technical documentation. Accessed: Oct. 25, 2026. [Online]. Available: https://open-meteo.com/en/docs/single-runs-api',
 '[5] FIMCM, ESPOL, “Análisis meteorológico y oceanográfico El impacto de las inusuales precipitaciones de octubre de 2026 en Guayaquil,” 2026. [Online]. Available: https://www.fimcm.espol.edu.ec/es/noticias/analisis-meteorologico-y-oceanografico-el-impacto-de-las-inusuales-precipitaciones-de',
 '[6] Ecuavisa, “68 sectores afectados por lluvias en Guayaquil,” Oct. 25, 2026. [Online]. Available: https://www.ecuavisa.com/amp/guayaquil/68-sectores-afectados-lluvias-guayaquil-tormenta-electrica-20261025-0011.html',
 '[7] “Zonas Inundables GYE,” ArcGIS FeatureServer, layer 0, effective date not supplied. Accessed: Oct. 25, 2026. [Online]. Available: https://services7.arcgis.com/NWWHhu45fOJtCgG3/arcgis/rest/services/Zonas_Inundables_GYE/FeatureServer/0',
 '[8] D. I. MacKenzie et al., “Estimating site occupancy rates when detection probabilities are less than one,” Ecology, vol. 83, no. 8, pp. 2248–2255, 2002, doi: 10.1890/0012-9658(2002)083[2248:ESORWD]2.0.CO;2.',
 '[9] IEEE, “Manuscript templates for conference proceedings.” Accessed: Oct. 25, 2026. [Online]. Available: https://cai.ieee.org/2025/manuscript-templates-for-conference-proceedings/',
 '[10] Prefectura del Guayas, “Samborondón.” Accessed: Oct. 25, 2026. [Online]. Available: https://guayas.gob.ec/cantones-2/samborondon/'
]

def base(paper=False):
 d=Document();s=d.sections[0]
 s.page_width=Cm(21);s.page_height=Cm(29.7)
 s.top_margin=Cm(1.9);s.bottom_margin=Cm(2.2 if paper else 2)
 s.left_margin=s.right_margin=Cm(1.65 if paper else 2.15)
 normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
 normal.font.color.rgb=RGBColor(0,0,0)
 normal.paragraph_format.space_after=Pt(5 if paper else 7)
 normal.paragraph_format.line_spacing=1.04 if paper else 1.1
 for name,size in [('Title',22 if paper else 25),('Subtitle',11),('Heading 1',11 if paper else 15),('Heading 2',10 if paper else 12)]:
  st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
  st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(5)
  if name.startswith('Heading'):st.font.bold=True
 d.core_properties.author='; '.join(AUTHORS)
 d.core_properties.title='FLOODSHADOW 3Jota'
 d.core_properties.subject='Hack2Impact 2026 Ecuador challenge'
 for st in d.styles:
  if st.type==1:
   st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
   for border in list(st.element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)
   for fonts in st.element.findall('.//'+qn('w:rFonts')):
    for key in list(fonts.attrib):
     if key.endswith('Theme'):del fonts.attrib[key]
 return d
def p(d,text,style=None,bold=False):
 x=d.add_paragraph(style=style);r=x.add_run(text);r.bold=bold
 r.font.name='Times New Roman';r.font.color.rgb=RGBColor(0,0,0)
 x.paragraph_format.widow_control=True
 return x
def h(d,text): return d.add_heading(text,level=1)
def table(d,headers,rows,paper=False):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 width=17.7 if paper else 16.7
 for column in t.columns:column.width=Cm(width/len(headers))
 for row in t.rows:
  for cell in row.cells:cell.width=Cm(width/len(headers))
 for c,x in zip(t.rows[0].cells,headers):c.text=str(x)
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr()
  if ri==0:trPr.append(OxmlElement('w:tblHeader'))
  trPr.append(OxmlElement('w:cantSplit'))
  for cell in row.cells:
   cell.width=Cm(width/len(headers))
   tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDE8EF' if ri==0 else ('F4F7F9' if ri%2==0 else 'FFFFFF'));tcPr.append(shade)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margins.append(e)
   tcPr.append(margins)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(2)
    for r in para.runs:r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=ri==0
 p(d,'')
 return t
def refs(d,heading=True):
 order=[]
 paragraphs=list(d.element.body.iter(qn('w:p')))
 for para in paragraphs:
  text=''.join(n.text or '' for n in para.iter(qn('w:t')))
  for value in re.findall(r'\[(\d+)\]',text):
   number=int(value)
   if 1<=number<=len(REFS) and number not in order:order.append(number)
 mapping={value:i+1 for i,value in enumerate(order)}
 for para in paragraphs:
  for node in para.iter(qn('w:t')):
   node.text=re.sub(r'\[(\d+)\]',lambda m:'['+str(mapping.get(int(m[1]),int(m[1])))+']',node.text or '')
 if heading:h(d,'References')
 for old in order:
  ref=re.sub(r'^\[\d+\]',f'[{mapping[old]}]',REFS[old-1])
  x=p(d,ref.replace('/', '/\u200b'))
  x.paragraph_format.first_line_indent=Cm(-.5);x.paragraph_format.left_indent=Cm(.5)
  for r in x.runs:r.font.size=Pt(8)
def page(d,title):
 d.add_page_break();h(d,title)

paper=base(True)
p(paper,'FLOODSHADOW Rainfall Intelligence and Evidence Verification for Guayaquil','Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
p(paper,'3Jota',bold=True).alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in AUTHORS:p(paper,name).alignment=WD_ALIGN_PARAGRAPH.CENTER
p(paper,'Hack2Impact 2026  Ecuador challenge').alignment=WD_ALIGN_PARAGRAPH.CENTER
section=paper.add_section(WD_SECTION_START.CONTINUOUS)
cols=section._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'1')
p(paper,'Abstract—FLOODSHADOW connects weather forecasts, territorial susceptibility, documented impacts and verification for intense rainfall in Guayaquil. The implemented prototype preserves 6652 official census polygons and groups them into 100 named analytical areas using a municipal nomenclature dataset and parish fallbacks. Two documented October 2026 storms and a separately preserved prospective weather capture support a connected before, rainfall replay and impact workflow. Observed station totals, model estimates, forecast outputs and candidate incident coordinates remain distinct. An explained index identifies places where evidence collection may be useful; it is not a calibrated flood probability. A separate Bayesian laboratory studies imperfect reporting and adaptive verification. Structural validation includes 38 passing automated tests, geometry checks and an interactive named-area filter. Real-event predictive effectiveness is not demonstrated: none of six calculable candidate sectors in the October 24–25 collection exceeds the exploratory priority threshold. The contribution is an auditable decision-support workflow and a transparent path toward field validation, rather than a claim of operational flood prediction.',bold=True)
p(paper,'Index Terms—rainfall, geographic information systems, underreporting, evidence provenance, active verification.',bold=True)
h(paper,'I  Problem and challenge alignment')
p(paper,'The Ecuador challenge asks how intelligent solutions can help Guayaquil anticipate and respond better to intense rain [1]. Evidence is fragmented across weather products, territorial layers and incident lists. A report at an intersection does not establish that its entire neighborhood flooded. Equally, no report does not establish that a place was unaffected. Conflating these situations can make an attractive map misleading.')
p(paper,'Our proposed users are analysts, researchers and staff reviewing where to collect evidence or investigate a rainfall episode. We support maps, data visualization, explainable prioritization and verification. The prototype does not issue official alerts, coordinate rescue operations, recommend safe routes or estimate economic losses. These functions require evidence and operational authority that the present dataset does not provide.')
h(paper,'II  Implemented solution')
p(paper,'An event identifier connects expected rainfall, hourly replay, later records and accepted verifications. The English interface uses Storm summary, Before rain, Rainfall replay, Reported impacts and Research lab. A persistent area search isolates both map geometry and analytical denominators. A selection remains available when changing stages.')
p(paper,'Users can switch between original census polygons and named-area outlines. Urdesa, for example, groups 89 complete polygons. Searching San Borondón resolves the Samborondón alias and shows its 246 INEC polygons separately. No Guayaquil forecast, priority or impact estimate is copied into that neighboring canton [2], [10].')
table(paper,['View','Question answered'],[
 ['Storm summary','What evidence exists for this storm?'],['Before rain','What rainfall did the earlier model run expect?'],['Rainfall replay','How did hourly model estimates change?'],['Reported impacts','What locations were reported and what remains unverified?'],['Research lab','How can imperfect reporting and verification affect inference?']],True)
p(paper,'Color legends explain meaning and units. Amber indicates relative watch priority. Blue indicates rainfall in mm or mm/h; its visible numerical scale follows the selected window. Red dots are provisional locations of published records. Gray means a missing value or lack of whole-sector verification, depending on the selected layer. Positive and negative sector-wide checks have separate states. The impact view starts with rainfall and report points rather than an entirely unverified gray surface.')
h(paper,'III  Data and geographic integrity')
p(paper,'INEC 2022 supplies 6652 unique canton-sector identifiers: 6416 under main-parish code 090150 and 236 rural polygons [2]. The main-parish code includes peripheral land and is not asserted to be a legal urban limit. Municipal nomenclature supplies 95 source polygons and recognizable place names [3]. A census polygon is assigned only when one named source area covers more than half of its original surface; remaining units use parish or peripheral groups. The resulting 100 groups form a complete, nonduplicated partition of the analyzed census units. Their outlines are analytical group boundaries, not a new cadastral delineation.')
p(paper,'Geographic calculations use WGS84 geodesic area and GEOS intersections. The original source had 7297 overlap pairs larger than 1 m². A declared deterministic partition removes analytical overlaps while preserving source files. Changes over 1% of original sector area exclude the territorial index in 897 sectors. Twelve invalid simplifications fall back to valid analytical geometry. All 6652 rendered sector geometries and 737 technical zone aggregates are valid. Membership, method, source URL and source checksum are retained.')
p(paper,'The documented susceptibility dataset contains 5746 polygons [7]. Their union is intersected with each analytical sector to prevent double counting. The effective date is unknown, so the footprint cannot prove that a location flooded in a particular past event. The interface describes it as susceptible land, not an incident history.')
p(paper,'Weather products use six distinct sampled ECMWF IFS cells per episode, with nominal resolution about 9 km [4]. A cell within 12 km of a representative sector point can be assigned; this does not increase spatial precision. Rain totals require complete, unique hourly timestamps. Missing values remain missing. ESPOL published five differing station totals for October 19, ranging from 112 to 181 mm [5]; exact instrument coordinates and accumulation windows are unavailable, so we neither interpolate these totals nor calculate a station-based forecast error.')
p(paper,'For October 24–25, the attributed Ecuavisa list contains 68 urban water-accumulation locations and ten other typed records, including three reported floods in Tenguel [6]. Eleven intersections have candidate coordinates. Only seven have geocoder alternatives consistent with one original sector; four remain sector-ambiguous. The remaining 67 records have no assigned coordinates. Point records never fill an entire polygon as confirmed flooding. Direct municipal infrastructure queries returned HTTP 500; drainage capacity, maintenance, obstructions, pumping and tides remain unknown.')
h(paper,'IV  Analysis and verification method')
p(paper,'The real-event vigilance index is a multicriteria convention. Let h be the fraction of sector land intersecting the susceptibility union. Let r be the midrank percentile of forecast accumulation among unique sampled cells: r = (Nless + 0.5 Nequal) / N. Dry forecasts contribute zero rain points. The index is')
p(paper,'I = 50h + 50r.                                      (1)',bold=True)
p(paper,'Both inputs must exist; missing factors are not silently reweighted. Equal importance and the threshold I ≥ 60 are explicit choices, not parameters fitted to the listed incidents. No calibrated flood probability or statistical confidence interval is claimed. Shadow Detection selects high-index sectors without sufficient localized evidence for further investigation.')
p(paper,'The original research laboratory separately models impact occurrence and reporting conditional on impact, with spatial latent fields and Poisson reporting, including false reports. Its design follows the general need to separate a hidden state from imperfect observation [8], but it is not a validated transfer of an ecological model to Guayaquil. Positive or negative experimental inspections reweight shared latent hypotheses and update other sectors. Priors and synthetic experiments remain explicitly experimental.')
p(paper,'Actual verification requires a storm, official sector, observation time, observer, method, descriptive evidence, HTTP(S) source and scope. Point observations do not count as whole-sector inspections. A negative sector-wide result needs representative evidence for the entire completed storm window. Entries start pending, can be accepted by the local operator and may be corrected without erasing history. New accepted evidence changes states, metrics and the verification queue; it does not rewrite the original assessment or automatically train a useful predictor from a few examples.')
p(paper,'Historical forecast runs were initialized before the selected storms, but the platform retrieved them later; reconstructed index values are not counted as predictions issued beforehand. A separate capture was received on October 25 at 18:34 UTC−5 for the window October 25 at 19:00 through October 28 at 19:00. Its inputs, index and SHA-256 checksum are preserved before that window. The provider’s exact issue time is unavailable. A checksum proves content consistency, not independently certified time or authenticity.')
h(paper,'V  Validation and observed results')
table(paper,['Check or metric','Observed result'],[
 ['Automated tests','38 passed; no failures'],['Census membership','6652 units, one group each'],['Urdesa filter','89 polygons; 3 mapped point records'],['October 24–25 comparison','0 of 6 calculable candidate sectors at I ≥ 60; 1 missing index'],['Real sector-wide verification','0 accepted; affected share not evaluable'],['SQL integrity and foreign keys','Passed'],['Production build','Passed; MapLibre size warning remains']],True)
p(paper,'The city’s model-assignment coverage is 99.44% of main-parish sectors but only 48.77% of their summed area. These denominators differ and neither is instrumental observation coverage. In the October 24–25 collection, seven sectors contain provisionally mapped points; this is not an estimate of the actual affected share of Guayaquil.')
p(paper,'The unfavorable 0/6 exploratory index comparison is reported without changing weights or threshold. It is not forecast recall: the index was reconstructed later, point locations remain provisional and the report collection is incomplete. Unreported units are never used as negative ground truth. Real-event ranking superiority and forecast calibration remain unproven.')
p(paper,'The preserved laboratory benchmark compared random selection, report counts, historical susceptibility, the original heuristic and the Bayesian model under the same candidate pools and inspection budget. The existing suite contains 2500 synthetic episodes. For hidden discoveries in its uneven-reporting setting, mean discoveries per five inspections were 0.980 for random/report-count selection, 1.632 for history, 1.636 for the original heuristic and 1.598 for the Bayesian model. Thus the model does not establish superiority over history in that setting. These are prior reproducible experimental results, not newly observed field outcomes; manifests and paired uncertainty estimates remain in docs/INFORME-v03.md.')
h(paper,'VI  Expected impact and implementation plan')
p(paper,'The intended impact is better evidence retrieval, clearer recognition of information gaps and a traceable basis for deciding what to inspect next. We have not measured reduced losses, faster emergency response or lives saved. Evaluation should measure time to find a source, correct interpretation of colors, coverage of representative inspections and discoveries with a demonstrably complete initial report collection.')
table(paper,['Phase','Next acceptance condition'],[
 ['Data stewardship','Confirm reuse terms, data owners and update cadence'],['Field protocol','Obtain representative positive and negative checks with compatible time windows'],['Prospective evaluation','Preserve forecasts before independent storms and compare fixed baselines'],['Model development','Fit only on training storms; assess calibration and held-out performance'],['Deployment','Shared database, authenticated review, monitoring and accessible usability tests']],True)
p(paper,'The present application runs with React, Vite, MapLibre GL JS and local JSON snapshots. Python, Shapely/GEOS, pyproj and pyshp prepare geographic data; SQLite provides a portable query copy. Verification records currently use browser storage and JSON export. Development used AI-assisted coding with Codex; no generative model fabricates observations, forecasts or incident classifications. Members must understand and defend the implementation. External sources retain attribution; redistribution terms must be checked for each complete source dataset.')
h(paper,'VII  Conclusion')
p(paper,'FLOODSHADOW delivers a connected and inspectable rainfall-evidence workflow with named geographic filtering and explicit uncertainty. It answers the Guayaquil challenge through maps, explainable prioritization and verification. Its current defensible result is functional integration with reproducible structural checks. The next scientific milestone is representative field evidence and independent prospective evaluation, not a stronger claim based on the present limited labels.')
refs(paper)
paper.save(OUT/'FLOODSHADOW_3Jota_Proposal_EN.docx')

guide=base(False)
p(guide,'FLOODSHADOW Guía para comprender y presentar el sistema','Title')
p(guide,'Equipo 3Jota',bold=True)
for name in AUTHORS:p(guide,name)
p(guide,'Hack2Impact 2026  Reto de Ecuador  25 de octubre de 2026')
h(guide,'I  Qué hace el proyecto')
p(guide,'FLOODSHADOW reúne lo que un modelo meteorológico esperaba, cómo cambió la lluvia estimada y qué afectaciones se publicaron. Permite explorar una zona concreta de Guayaquil y reconocer qué información falta. El usuario puede añadir evidencia real y revisar cómo cambia el estado de un sector. Su valor actual está en conectar datos, explicar sus límites y orientar verificaciones; todavía no demuestra que prediga inundaciones con precisión.')
p(guide,'La aplicación está en inglés por el jurado internacional. Esta guía explica sus términos en español. La propuesta que debe entregarse está en inglés y se incluye como documento separado. Las referencias se numeran en orden IEEE. El diseño y la lectura tienen prioridad sobre imitar una publicación certificada por IEEE [9].')
page(guide,'II  Qué significa cada color')
table(guide,['Color y capa','Lectura correcta'],[
 ['Ámbar en Watch priority','Prioridad relativa para prestar atención o buscar evidencia. 0–29 menor, 30–59 moderada, 60–100 mayor. No es porcentaje de inundación.'],['Azul en Expected rain o Rain estimate','Cantidad de lluvia, en mm o mm/h según la línea de tiempo. Más oscuro significa más lluvia dentro de la escala numérica visible.'],['Punto rojo','Lugar candidato de un registro publicado. Puede necesitar revisión geográfica y no delimita un barrio inundado.'],['Gris en lluvia o prioridad','No existe un valor utilizable para esa capa. No significa que sea seguro ni que el valor sea cero.'],['Gris en Verification status','No existe verificación aceptada de todo el polígono. Puede haber un reporte puntual sin inspección sectorial.'],['Verde o rojo en Verification status','Resultado sectorial negativo o positivo aceptado, con alcance representativo. No se precargaron verificaciones ficticias.']])
p(guide,'Un área puede verse del mismo azul porque sus polígonos comparten una celda meteorológica de aproximadamente 9 km. El mapa no inventa variación calle por calle. La reproducción horaria sí muestra cómo cambia el valor del modelo con el tiempo.')

page(guide,'III  Buscar una zona y sus subdivisiones')
p(guide,'Escribe Urdesa en Find a named area y selecciona el resultado. El mapa deja solo sus 89 polígonos, las cifras usan ese mismo conjunto y la ficha de área muestra únicamente los registros localizados allí. La fuente municipal contiene 95 áreas de nomenclatura. La asignación a los polígonos INEC produce 100 grupos, incluidos grupos de periferia y parroquias rurales [2], [3].')
p(guide,'Census polygons conserva las unidades pequeñas existentes. Named areas agrega contornos y nombres de las agrupaciones. Un grupo puede contener muchos polígonos, pero no los recorta ni cambia sus identificadores oficiales. Al abrir uno aparece un nombre como Urdesa junto al identificador; el código completo permanece disponible para trazabilidad.')
p(guide,'Para asignar el nombre se compara la superficie original del polígono con el área municipal. Se requiere que una sola área cubra más de la mitad. Si esto no puede establecerse, se conserva una agrupación de parroquia o periferia. El contorno resultante es una unidad de consulta y no un nuevo límite legal. La página municipal y el método quedan documentados.')
p(guide,'El buscador reconoce Samborondón y la variante San Borondón. Ese nombre pertenece a otro cantón [10]. Se muestran 246 unidades INEC como contexto geográfico, sin copiarles la lluvia, los incidentes o el índice de Guayaquil. Para continuar el análisis real del episodio, elige una zona dentro de la cobertura analizada.')
h(guide,'Ejemplo comprobable de Urdesa')
p(guide,'En el episodio 24–25 de octubre, Urdesa reúne 89 polígonos, tres registros puntuales localizados y un valor de lluvia modelada de 22,3 mm compartido por sus celdas asignadas. La unión susceptible intersecta aproximadamente 1,05 km² de 2,95 km² con geometría histórica utilizable. Esa superficie expresa susceptibilidad documentada, no el área que se inundó ese día.')
p(guide,'Los registros sin coordenadas se conservan en el catálogo general; no se fuerzan dentro de una agrupación. Una mención textual a una ciudadela no basta para atribuirle un polígono preciso. Por ello, los números filtrados pueden ser menores que el total mencionado por una noticia.')

page(guide,'IV  Entender las pantallas')
table(guide,['Pantalla en inglés','Qué consultar'],[
 ['Storm summary','Resumen del episodio. Observa lugares publicados, cobertura, candidatos para evidencia y el significado de la capa elegida.'],['Before rain','Lo que esperaba un modelo de lluvia para ese episodio. Consulta milímetros previstos y las contribuciones de la prioridad territorial.'],['Rainfall replay','Recorrido de estimaciones horarias del modelo. Usa Play rainfall replay, el deslizador o las flechas. No es un radar ni un seguimiento instrumental en vivo.'],['Reported impacts','Lluvia estimada después y ubicaciones reportadas. Revisa verificaciones, denominadores y comparación con las condiciones esperadas.'],['Research lab','Experimento separado para estudiar subregistro y verificar sectores sintéticos. Esos incidentes no describen la ciudad actual.']])
h(guide,'Qué quiere decir un pronóstico guardado de un episodio pasado')
p(guide,'Un pronóstico es lo que el modelo esperaba que lloviera. Para reconstruir octubre se recuperó después una corrida que se había inicializado antes del episodio [4]. Eso permite examinar la expectativa meteorológica, pero no demuestra que FLOODSHADOW publicara ese análisis antes de la lluvia. Inicialización del modelo, publicación del proveedor y recepción por la aplicación son tiempos diferentes.')
p(guide,'Además existe una captura recibida el 25 de octubre a las 18:34, antes de su ventana de 72 horas iniciada a las 19:00, en hora de Guayaquil. Esa evaluación queda preservada para compararla con observaciones futuras. No se debe convertir la reconstrucción histórica en un acierto previo ni reemplazar valores iniciales después de ver los resultados.')
h(guide,'Recorrido de cinco acciones')
p(guide,'1. Selecciona el episodio 24–25 de octubre y busca Urdesa. 2. Abre Before rain y un polígono para leer lluvia esperada y prioridad. 3. Cambia a Rainfall replay y reproduce las horas. 4. Abre Reported impacts para contrastar lluvia estimada y registros, manteniendo el mismo lugar. 5. Consulta la fuente o registra una verificación real; usa Research lab cuando necesites ensayar resultados inventados.')

page(guide,'V  Datos y método sin confusiones')
p(guide,'La geografía analizada tiene 6652 polígonos cantonales, con 6416 de cabecera y 236 rurales. El código de cabecera incluye periferia; no se redefine con él un límite urbano legal. La asignación meteorológica por punto representativo alcanza al 99,44% de los polígonos de cabecera, pero al 48,77% de su superficie sumada. Son métricas distintas y no miden cobertura instrumental.')
p(guide,'ESPOL publicó cinco acumulados del 19 de octubre: Cerro Santa Ana 181 mm, Cisne II 141 mm, Martha de Roldós 132 mm, ESPOL 117 mm y Aeropuerto 112 mm [5]. Al no existir coordenadas instrumentales y ventanas exactas compatibles, el sistema no inventa una distribución horaria para esas estaciones ni las interpola por barrios.')
p(guide,'La colección del 24–25 contiene 78 registros de distintos tipos. Incluye 68 acumulaciones urbanas y otras diez novedades; once cruces tienen coordenadas candidatas, siete una asignación sectorial provisional consistente y cuatro sector ambiguo. Los otros 67 registros no reciben coordenadas [6]. No son 68 barrios totalmente inundados.')
p(guide,'La prioridad se calcula como I = 50h + 50r. h es la fracción de superficie que intersecta la unión de susceptibilidad; r es el rango percentil de lluvia prevista entre seis celdas únicas. La lluvia cero aporta cero puntos. Ambos factores deben existir. Los pesos y el umbral 60 son convenciones explícitas, no parámetros calibrados con inundaciones reales.')
p(guide,'La huella histórica contiene 5746 geometrías [7]. No documenta fechas de inundación por barrio. Cuando el usuario solicita historial territorial, se muestra lo que sí puede respaldarse: superficie susceptible, fuente y sus límites. La capacidad hidráulica, mantenimiento y obstrucción de drenajes, topografía detallada y marea siguen sin datos integrados. La consulta de infraestructura municipal falló con HTTP 500; la nueva nomenclatura sí pudo descargarse por otro canal oficial.')
h(guide,'Qué se validó y qué sigue pendiente')
p(guide,'Se comprobaron IDs, pertenencia única por grupo, geometrías, horas, faltantes, cálculos, correcciones y trazabilidad. Hay 38 pruebas automatizadas aprobadas. Eso valida estructura y comportamiento del software. No valida en campo la ocurrencia de inundaciones ni calibra una probabilidad. En el caso 24–25, ninguno de seis polígonos candidatos con índice calculable supera 60; un séptimo carece de índice. No se ocultó ese resultado ni se bajó el umbral para aparentar éxito.')

page(guide,'VI  Verificación y aprendizaje')
p(guide,'Una verificación añade evidencia observada, no una opinión basada en el color. Debe indicar episodio, polígono, momento, resultado, responsable, método, descripción, enlace y alcance. Se guarda pendiente y requiere revisión declarada por el operador. La aplicación todavía no autentica a una entidad inspectora ni sincroniza varios dispositivos.')
p(guide,'Una inspección puntual no confirma un polígono completo. Para un resultado sectorial negativo se necesita evidencia representativa de toda la ventana terminada; una visita cuando ya dejó de llover no demuestra que nunca hubo afectación. Las correcciones conservan el registro anterior. La evidencia aceptada más reciente cambia el estado, las métricas y la siguiente cola de revisión.')
p(guide,'Los datos reales no reentrenan automáticamente un predictor útil con tan pocas muestras. El laboratorio demuestra otra función: una observación experimental positiva o negativa actualiza inferencias espaciales. Separar ocurrencia de observación se relaciona con métodos de detección imperfecta [8], pero no establece que el modelo esté calibrado para Guayaquil.')
h(guide,'Respuestas claras para el jurado')
table(guide,['Pregunta','Respuesta defendible'],[
 ['¿El 80 significa 80% de inundación?','No. Es una prioridad relativa de 80 sobre 100 según dos factores disponibles.'],['¿Por qué varios polígonos tienen igual azul?','Comparten una celda meteorológica gruesa; no se inventa precisión local.'],['¿Gris significa que no pasó nada?','No. Significa dato faltante o ausencia de verificación completa, según la capa.'],['¿Aprende con una inspección?','Actualiza evidencia y recomendaciones; el aprendizaje probabilístico se demuestra en el laboratorio.'],['¿Supera al algoritmo anterior?','No se demuestra superioridad real. El benchmark sintético conserva comparadores y resultados mixtos.'],['¿Qué falta para operar?','Evidencia representativa, validación prospectiva, fuentes compatibles y revisión institucional.']])
h(guide,'Resultados comparativos del laboratorio')
p(guide,'La evaluación conservada usa 2500 episodios sintéticos y cinco comparadores bajo las mismas condiciones. En Uneven reporting, los descubrimientos ocultos medios por cinco inspecciones son 0,980 con azar o reportes, 1,632 con antecedentes, 1,636 con heurística original y 1,598 con el modelo. Ese escenario no demuestra superioridad sobre antecedentes o heurística. Las simulaciones no son inspecciones reales. Los manifiestos y los intervalos completos permanecen en INFORME-v03.md.')

page(guide,'VII  Guion del video en inglés')
p(guide,'Video obligatorio máximo de tres minutos. El guion siguiente está pensado para unas 2 min 40 s a 3 min, según pausas y velocidad. Hay que cronometrar una grabación real; contar palabras no certifica duración. Muestra la aplicación mientras se pronuncia el texto.')
pitch='''Hello, we are 3Jota: José Luis Paladines Sánchez, John Maximiliano Carreño Robalino and Julinho Jesús Plua Mendosa. We built FLOODSHADOW for the Ecuador challenge: helping Guayaquil anticipate and respond better to intense rainfall.

The problem is not simply a missing map. Weather forecasts, territorial information and incident reports describe different things. A report at one intersection does not mean an entire neighborhood flooded. A place without reports is not necessarily safe.

FLOODSHADOW connects this evidence around one storm. We preserve 6652 official census polygons and organize them into recognizable named areas. Here, searching Urdesa isolates 89 polygons and updates the map, statistics and evidence together.

Before rain shows what an earlier weather-model run expected. Amber means relative watch priority, not flood probability. Rainfall replay shows hourly model estimates. Darker blue means more rain on the displayed scale. Nearby polygons may share a coarse weather cell, so we do not invent street-level precision.

Reported impacts connects modeled rainfall with published locations. Red points are candidate locations, not flooded neighborhood boundaries. Our October 24–25 collection contains 78 typed records. Ambiguous places remain unresolved instead of receiving invented coordinates.

Shadow Detection asks where further evidence may be useful. A real verification records the source, time, method and scope, then enters review. Accepted evidence updates status and the next verification queue. Our separate Research lab demonstrates how experimental positive and negative checks update spatial inferences.

We are transparent about limitations. The exploratory index did not highlight any of the six calculable candidate sectors in this historical comparison. We do not claim validated flood prediction. The software passes 38 automated tests, while field validation remains a next step.

Our contribution is a usable, traceable workflow: understand expected rainfall, inspect documented impacts, identify missing evidence and learn from valid checks. FLOODSHADOW helps the user decide what to investigate next, while keeping uncertainty visible. Thank you.'''
for block in pitch.split('\n\n'):p(guide,block)

page(guide,'VIII  Presentación en vivo y entrega')
p(guide,'El pitch final previsto por las reglas dura de cinco a siete minutos y también debe ser en inglés [1]. Usa este recorrido con tiempos orientativos: 0:00–0:45 problema y usuario; 0:45–1:30 selección de Urdesa; 1:30–2:20 Before rain; 2:20–3:10 Rainfall replay; 3:10–4:00 Reported impacts; 4:00–5:00 verificación y Research lab; 5:00–6:00 resultados reales, límites y siguiente validación. No presentes una inspección sintética como evidencia de campo.')
table(guide,['Requisito de la convocatoria','Cobertura de la entrega'],[
 ['Propuesta PDF en inglés','FLOODSHADOW_3Jota_Proposal_EN.pdf incluye problema, solución, tecnología, impacto esperado y plan.'],['Video en inglés de hasta 3 minutos','Guion incluido. El equipo debe grabar, cronometrar y entregar el archivo de video.'],['Repositorio opcional','Código y datos locales disponibles. No se inventó un enlace público ni se publicó el proyecto.'],['Un desafío por equipo','Reto Ecuador centrado en Guayaquil. Samborondón aparece como contexto separado.'],['Herramientas y uso de IA declarados','React, Vite, MapLibre, Python, Shapely, pyproj, pyshp, SQLite, fuentes públicas y asistencia de Codex.'],['Archivos accesibles sin permisos especiales','PDF local sin contraseña. Verificar accesibilidad de los enlaces que el equipo entregue.'],['Equipo de 2 a 3 personas','Tres autores identificados como 3Jota. La inscripción y afiliación no se verifican desde el código.']])
p(guide,'El documento de reglas suministrado establece entrega el lunes 26 de octubre de 2026 hasta las 21:00 UTC−5 y pitch final el martes 27 de 09:00 a 10:00, con conexión diez minutos antes. El usuario comunicó que se eliminaron las restricciones sobre trabajo previo y comienzo del historial; esta guía no las vuelve a imponer ni afirma haber verificado un comunicado actualizado del comité. Se debe usar el canal de entrega que el comité comunique.')
h(guide,'Ejecutar y comprobar')
p(guide,'Desde la carpeta del proyecto: npm install, npm run dev y abrir http://127.0.0.1:5174/. Ejecutar npm test y npm run build. No se requieren claves API para leer los snapshots incluidos. El mapa base necesita internet y WebGL; los datos locales conservan su procedencia. Las verificaciones se guardan en ese navegador y pueden exportarse en JSON.')
p(guide,'Antes de entregar: abrir el PDF final, revisar nombres y fuentes, grabar el video en inglés, medir su duración real, comprobar permisos del enlace y realizar una sola entrega por equipo. Los documentos no sustituyen la inscripción, la grabación, la validación de campo ni la aprobación del comité.')
page(guide,'IX  Referencias')
refs(guide,False)
guide.save(OUT/'FLOODSHADOW_3Jota_Guia_ES.docx')
print('Created two documents in',OUT)
print('Video script words:',len(pitch.split()))
