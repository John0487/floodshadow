"""Plain-language presentation for 3Jota. No IEEE manuscript styling."""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'docs'/'entrega';OUT.mkdir(parents=True,exist_ok=True)
d=Document();section=d.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7)
section.left_margin=section.right_margin=Cm(2);section.top_margin=Cm(1.7);section.bottom_margin=Cm(1.8)
for s in d.styles:
 if s.type==1:
  s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
  for b in list(s.element.findall('.//'+qn('w:pBdr'))):b.getparent().remove(b)
  for fonts in s.element.findall('.//'+qn('w:rFonts')):
   for key in list(fonts.attrib):
    if key.endswith('Theme'):del fonts.attrib[key]
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.line_spacing=1.12
d.styles['Normal'].paragraph_format.space_after=Pt(8)
for name,size in [('Title',36),('Subtitle',17),('Heading 1',23),('Heading 2',14)]:
 d.styles[name].font.size=Pt(size);d.styles[name].font.bold=name!='Subtitle';d.styles[name].paragraph_format.space_before=Pt(10);d.styles[name].paragraph_format.space_after=Pt(9)
d.core_properties.author='José Luis Paladines Sánchez; John Maximiliano Carreño Robalino; Julinho Jesús Plua Mendosa'
d.core_properties.title='FLOODSHADOW A clearer picture of every storm'
def p(text,style=None,bold=False):
 x=d.add_paragraph(style=style);r=x.add_run(text);r.font.name='Arial';r.bold=bold;x.paragraph_format.widow_control=True;return x
def h(text):d.add_heading(text,level=1)
def sub(text):d.add_heading(text,level=2)
def page(title,lead):d.add_page_break();h(title);p(lead)
def table(headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 widths=widths or [17/len(headers)]*len(headers)
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for c,text in zip(t.rows[0].cells,headers):c.text=text
 for values in rows:
  for c,text in zip(t.add_row().cells,values):c.text=str(text)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for c,w in zip(row.cells,widths):
   c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'173F56' if i==0 else 'EEF5F8' if i%2 else 'FFFFFF');pr.append(shade)
   edges=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    n=OxmlElement('w:'+edge);n.set(qn('w:val'),'single');n.set(qn('w:sz'),'4');n.set(qn('w:color'),'D9D9D9');edges.append(n)
   pr.append(edges);m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    n=OxmlElement('w:'+edge);n.set(qn('w:w'),'105');n.set(qn('w:type'),'dxa');m.append(n)
   pr.append(m)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(5)
    for r in para.runs:r.font.name='Arial';r.font.size=Pt(10);r.bold=i==0;r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
 p('')

p('FLOODSHADOW','Title')
p('A clearer picture of every storm','Subtitle')
p('Understand the rain. See the evidence. Know what needs a closer look.')
image=ROOT/'docs'/'screenshots'/'final-prepared-demo.png'
if image.exists():
 d.add_picture(str(image),width=Cm(17))
 x=p('Urdesa in the prepared local demonstration. Blue shows modeled rainfall; the red points are candidate locations from published reports.');x.runs[0].font.size=Pt(9)
p('3Jota  |  Hack2Impact 2026',bold=True)
p('José Luis Paladines Sánchez\nJohn Maximiliano Carreño Robalino\nJulinho Jesús Plua Mendosa')

page('The problem is a fragmented picture','Heavy rain affects different places in different ways. Yet a weather forecast, a flood-prone land map and a list of street incidents answer different questions.')
sub('What the user needs to know')
p('Where should we pay attention? What was reported? What still needs a check? FLOODSHADOW brings these questions together around one selected storm, with recognizable place names and links back to the evidence.')
table(['Before the Rain','During the Rain','After the Rain'],[['What rain was expected?','How did the available rain estimates change?','Where were impacts documented?'],['Review places that may deserve attention.','Replay the saved hourly evidence.','Read the findings and the remaining gaps.']])
sub('A useful next decision')
p('An analyst can choose where to investigate next and explain why. A presenter can follow one neighborhood through the same event instead of switching between unrelated maps.')
sub('What makes this approach different')
p('The map shows both evidence and its limits. A reported intersection is not expanded into an entirely flooded neighborhood. A place without reports is not marked safe. A later reconstruction is not presented as a forecast the platform issued before the storm.')
p('This addresses the Ecuador challenge: helping Guayaquil anticipate and respond better to heavy rainfall through clearer information, explainable monitoring and evidence collection.')

page('One neighborhood you can follow','Search Urdesa. The map, local figures and evidence panel all focus on the same selected area.')
table(['In the prepared October 24–25 example','What it means'],[['89 smaller areas','Official census polygons are kept together under the recognizable name Urdesa.'],['22.3 mm of modeled rainfall','The same coarse weather cell can cover many smaller polygons. This is rain accumulation, not street flood depth.'],['3 mapped records','Three candidate locations from published water-accumulation reports. They do not establish the extent of flooding.'],['1.05 km² of susceptible land','Intersection with a documented flood-prone footprint, within 2.95 km² with usable historical geometry. This is not the area flooded in this event.']],[6.7,10.3])
sub('Start broad and look closer')
p('The prepared geography covers 6,652 Guayaquil census polygons in 100 named analytical groups. Municipal quadrants provide general-area navigation. Smaller technical units remain available when a local boundary has no verified neighborhood name.')
sub('Keep neighboring places distinct')
p('Searching San Borondón finds Samborondón as a separate canton. Its 246 census boundaries provide context, but Guayaquil’s rainfall estimates and impact records are not copied across the river.')
p('Whole-polygon group outlines support navigation. They are not new legal neighborhood boundaries.')

page('Read a finding in a few seconds','Every finding tells you the place, the reason, the evidence, the time context and what remains uncertain. Findings are informational; they are not official emergency alerts.')
table(['Finding','What the user should understand'],[['Preventive Watch','The rain forecast and land context suggest closer attention. No flood is confirmed.'],['Documented Impact','A published source reports an incident at a candidate location. Its exact location and extent may still need review.'],['Confirmed Impact','A reviewed operator check documents an impact within its stated scope. A point check does not confirm an entire neighborhood.'],['Verification Needed','Conditions justify a closer look, but the available evidence does not establish what happened.'],['Data Unavailable','An essential input is missing. Unknown does not mean safe.']],[5.3,11.7])
sub('New evidence has a history')
p('A real check records who observed what, when, how and where the evidence can be found. It starts pending review. Accepted evidence changes the assessment and the next verification queue. A correction preserves the earlier entry instead of erasing it.')
sub('Colors explain a specific layer')
p('Amber shows relative watch priority. Darker blue means more rain on the displayed scale. Red points identify published locations. Gray means an unusable value or missing whole-area verification, according to the selected layer. Text and icons carry the meaning alongside color.')

page('Ask what might change','The Rainfall Scenario Explorer starts from a saved real forecast. The user changes an assumption while the original sources remain untouched.')
p('Simulation — Not a real-time observation',bold=True)
table(['Rain amount assumption','Smaller areas at the watch threshold'],[['0% of the saved forecast','0 of 6,416'],['100% of the saved forecast','191 of 6,416'],['200% of the saved forecast','1,233 of 6,416']],[8,9])
p('These are checked sensitivity results for the prepared October 24–25 main-parish selection, with duration unchanged, no north/south tilt and all sampled cells retained. They are not counts of flooded places or forecasts of damage.')
sub('Explore one assumption at a time')
p('Increase or decrease rain, extend the assumed duration, tilt the spatial pattern or remove weather-cell availability. The map recalculates exploratory priority against the original rainfall reference. Missing inputs stay missing. No simulated incident enters the real evidence history.')
sub('Why this helps a demonstration')
p('The team can explain how assumptions influence a recommendation without waiting for another storm. The experiment also reveals what information would be needed before making a stronger claim.')
sub('Learning has three different meanings')
p('Collecting data saves evidence. Updating an assessment changes a finding when new evidence arrives. Retraining a model fits its parameters to independent examples. The first two are implemented for real evidence; automatic supervised retraining is a future capability.')

page('Ready to demonstrate locally','The prepared beta runs on one computer, with saved sources, a local map and a file-based evidence history. It does not depend on a weather provider answering during the presentation.')
table(['What is working','What still needs evidence'],[['42 passing automated tests and a production build','Field validation and independently reviewed outcomes'],['Named-area search, rainfall replay and informational findings','Compatible live station feeds and better location review'],['Separate scenarios and a downloadable impact PDF','Calibrated flood probabilities and prospective accuracy'],['Local saved checks, corrections and finding history','Authenticated institutional review and shared deployment']],[8.5,8.5])
sub('What we can defend today')
p('FLOODSHADOW makes rainfall evidence easier to inspect and explain. We have not measured reduced losses or faster emergency response. The historical index did not highlight any of six calculable candidate sectors in the October 24–25 comparison; one other candidate lacks an index. We show this result rather than adjusting the story to hide it.')
sub('The next step')
p('Collect representative positive and negative checks across independent storms. Preserve each earlier assessment, compare it with compatible later evidence and evaluate simple baselines before training a more complex model. The intended impact is better-informed monitoring and verification; its practical benefit still needs a user and field study.')
sub('Sources and technology')
p('Prepared sources include INEC 2022 geography, Guayaquil municipal nomenclature, ECMWF through Open-Meteo, ESPOL’s published station totals and Ecuavisa’s incident list attributed to Segura EP. Five published station totals and model estimates are kept separate. Drainage condition, detailed terrain, tides and compatible live station observations remain unavailable.')
p('The local beta uses React, Vite, MapLibre, JavaScript, Python geographic processing and structured files. Development used AI-assisted coding with Codex; a generative model does not create measurements or incidents. The team’s English manuals retain source links, technical detail and feature status.')
p('Team 3Jota  |  José Luis Paladines Sánchez  |  John Maximiliano Carreño Robalino  |  Julinho Jesús Plua Mendosa',bold=True)

for i,sec in enumerate(d.sections):
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=footer.add_run('FLOODSHADOW  ·  3Jota  ·  Hack2Impact 2026');r.font.name='Arial';r.font.size=Pt(8)
d.save(OUT/'FLOODSHADOW_3Jota_Presentation_EN.docx')
print('Created visual presentation')
